#!/usr/bin/env python
"""
Quick Test with Your Data - Step by Step

This script will help you test the updated processor with your actual WP LOA file
"""

from processors.wp_loa_formula import WPLOAFormulaProcessor
from pathlib import Path
import openpyxl

print("=" * 80)
print("WP LOA PROCESSOR - TESTING GUIDE")
print("=" * 80)

# Step 1: Load file
print("\n📋 Step 1: Loading your WP LOA Report file...")

input_file = None
uploads = Path('uploads')

# Check if file exists in uploads folder
if uploads.exists():
    xlsx_files = list(uploads.glob('*.xlsx')) + list(uploads.glob('*.xls'))
    if xlsx_files:
        input_file = xlsx_files[0]
        print(f"   ✓ Found: {input_file.name}")
    else:
        print("   ✗ No Excel files in 'uploads' folder")
else:
    print("   ℹ To use this test, place your file in 'uploads' folder")
    print("   Example:")
    print("      uploads/")
    print("      └── YourFileName.xlsx")
    exit()

if not input_file:
    print("\n   📝 Usage:")
    print("      1. Copy your WP LOA Report file to 'uploads' folder")
    print("      2. Run this script again")
    exit()

# Step 2: Process the file
print("\n📝 Step 2: Processing with UPDATED formulas...")
print("   - Extracting PID parts (K, L, M, N)")
print("   - Creating L1 formula (O)")
print("   - Creating VLOOKUP formulas (Q-U, Z-AA)")

try:
    processor = WPLOAFormulaProcessor(str(input_file))
    processor.load_file()
    processor.create_formulas()
    
    output_file = uploads / f"TEST_output_{input_file.name}"
    processor.save_workbook(str(output_file))
    print(f"   ✓ Processed and saved to: {output_file.name}")
except Exception as e:
    print(f"   ✗ Error: {str(e)}")
    import traceback
    traceback.print_exc()
    exit()

# Step 3: Verify the output
print("\n🔍 Step 3: Verifying output...")

try:
    # Load the output file to check formulas
    wb = openpyxl.load_workbook(str(output_file))
    ws = wb.active
    
    # Check row 2 (first data row)
    print("\n   📊 Checking Row 2 (first data row):")
    
    # Column K
    k_val = ws['K2'].value
    print(f"   • K2 (PID Part 1): {k_val}")
    if isinstance(k_val, str) and k_val.startswith('='):
        print(f"      ✓ Formula detected")
    
    # Column L
    l_val = ws['L2'].value
    print(f"   • L2 (PID Part 2): {l_val}")
    if isinstance(l_val, str) and l_val.startswith('='):
        print(f"      ✓ Formula detected")
    
    # Column M
    m_val = ws['M2'].value
    print(f"   • M2 (PID Part 3): {m_val}")
    if isinstance(m_val, str) and m_val.startswith('='):
        print(f"      ✓ Formula detected")
    
    # Column N
    n_val = ws['N2'].value
    print(f"   • N2 (PID Part 4): {n_val}")
    if isinstance(n_val, str) and n_val.startswith('='):
        print(f"      ✓ Formula detected")
    
    # Column O (L1)
    o_val = ws['O2'].value
    print(f"   • O2 (L1 formula): {o_val}")
    if o_val == '=L2&"-"&M2&"-"&N2':
        print(f"      ✓ Correct L1 formula!")
    
    # Column Q (First VLOOKUP)
    q_val = ws['Q2'].value
    print(f"   • Q2 (PROGRAM_MBR): {q_val}")
    if isinstance(q_val, str) and 'VLOOKUP(O2' in q_val:
        print(f"      ✓ VLOOKUP references O (L1) - CORRECT!")
    elif isinstance(q_val, str) and 'VLOOKUP(P2' in q_val:
        print(f"      ✗ VLOOKUP still references P (old version)")
    
    # Check header row
    print("\n   📋 Checking Headers (Row 1):")
    headers_ok = True
    expected = ['PID (Mother and Sub)', '1', 'YEAR', '3', 'L1', 'L2']
    for col, expected_header in zip('KLMNOP', expected):
        actual = ws[f'{col}1'].value
        match = "✓" if actual == expected_header else "✗"
        print(f"      {match} {col}1: {actual}")
        if actual != expected_header:
            headers_ok = False
    
    wb.close()
    
    # Step 4: Summary
    print("\n" + "=" * 80)
    print("✅ TESTING COMPLETE!")
    print("=" * 80)
    print(f"\nOutput file: {output_file.name}")
    print("\n📝 Next Steps:")
    print("   1. Open the output file in Excel")
    print("   2. Check columns K-N - should show extracted parts")
    print("   3. Check column O - should show L1 concatenation (e.g., 'BSRF-26-SA')")
    print("   4. Check columns Q-U - should show VLOOKUP results:")
    print("      • If showing values → VLOOKUP working! ✓")
    print("      • If showing 'N/A' → Check BUDGET sheet Column A format")
    print("      • If showing '#NAME?' → Check for external file issues")
    print("\n💡 TIPS:")
    print("   • Click on a cell to see the formula in the formula bar")
    print("   • VLOOKUP should reference O2, not P2")
    print("   • All formulas should start with =")
    
except Exception as e:
    print(f"   ✗ Error verifying output: {str(e)}")
    import traceback
    traceback.print_exc()
